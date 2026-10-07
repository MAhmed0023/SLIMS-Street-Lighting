# -*- coding: utf-8 -*-
"""Shared document scaffolding for the SLIMS BRD."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x1F, 0x4E, 0x79)
BLUE = RGBColor(0x2E, 0x74, 0xB5)
GREY = RGBColor(0x59, 0x59, 0x59)
HDR_BG = "1F4E79"
ALT_BG = "EFF4FA"

doc = Document()

# ------------------------------------------------------------------ styles
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.08

for name, size, color, bold in [("Heading 1", 17, NAVY, True),
                                ("Heading 2", 13.5, BLUE, True),
                                ("Heading 3", 11.5, NAVY, True),
                                ("Heading 4", 10.5, GREY, True)]:
    s = doc.styles[name]
    s.font.name = "Calibri"
    s.font.size = Pt(size)
    s.font.color.rgb = color
    s.font.bold = bold
    s.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
    s.paragraph_format.space_after = Pt(6)
    s.paragraph_format.keep_with_next = True

sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.0)
sec.top_margin = Cm(1.9)
sec.bottom_margin = Cm(1.9)


# ----------------------------------------------------------------- helpers
def shade(cell, hexcolor):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hexcolor)
    cell._tc.get_or_add_tcPr().append(el)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def H1(text, page_break=False):
    if page_break:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    p = doc.add_heading(text, level=1)
    pb = OxmlElement("w:pBdr")
    bt = OxmlElement("w:bottom")
    bt.set(qn("w:val"), "single")
    bt.set(qn("w:sz"), "10")
    bt.set(qn("w:space"), "3")
    bt.set(qn("w:color"), "2E74B5")
    pb.append(bt)
    p._p.get_or_add_pPr().append(pb)
    return p


def H2(text):
    return doc.add_heading(text, level=2)


def H3(text):
    return doc.add_heading(text, level=3)


def P(text="", bold=False, italic=False, size=None, align=None, space=6):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if align == "c":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "j":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(space)
    return p


def BUL(items, style="List Bullet"):
    for it in items:
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent = Cm(0.75)
        if isinstance(it, tuple):
            r = p.add_run(it[0] + " ")
            r.bold = True
            p.add_run(it[1])
        else:
            p.add_run(it)


def NUM(items):
    BUL(items, style="List Number")


def TBL(headers, rows, widths=None, font=9, hdr_font=9, align_center=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    hdr = t.rows[0]
    repeat_header(hdr)
    for i, h in enumerate(headers):
        c = hdr.cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(hdr_font)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        c.paragraphs[0].paragraph_format.space_after = Pt(2)
        c.paragraphs[0].paragraph_format.space_before = Pt(2)
        shade(c, HDR_BG)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            para = cells[i].paragraphs[0]
            para.paragraph_format.space_after = Pt(2)
            para.paragraph_format.space_before = Pt(2)
            txt = str(val)
            bold = txt.startswith("**")
            run = para.add_run(txt.replace("**", ""))
            run.font.size = Pt(font)
            run.bold = bold
            if align_center and i in align_center:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if ri % 2 == 1:
                shade(cells[i], ALT_BG)
    if widths:
        for r in t.rows:
            for i, w in enumerate(widths):
                r.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def CAPTION(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = GREY
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    return p


def footer_text(section, text):
    p = section.footer.paragraphs[0]
    p.text = ""
    r = p.add_run(text)
    r.font.size = Pt(8)
    r.font.color.rgb = GREY
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_page_numbers(section):
    p = section.footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for instr in ("PAGE",):
        f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
        it.text = " " + instr + " "
        f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
        run._r.append(f1); run._r.append(it); run._r.append(f2)
    run.font.size = Pt(8)
    run.font.color.rgb = GREY


def toc_field():
    p = doc.add_paragraph()
    run = p.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
    it.text = r' TOC \o "1-2" \h \z \u '
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t")
    t.text = "Right-click here and choose 'Update Field' to build the table of contents."
    f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "end")
    run._r.append(f1); run._r.append(it); run._r.append(f2)
    run._r.append(t); run._r.append(f3)
