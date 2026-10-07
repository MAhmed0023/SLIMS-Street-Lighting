# -*- coding: utf-8 -*-
"""BRD sections: cover through master data."""
from docx_kit import *
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# =========================================================== COVER PAGE
for _ in range(3):
    doc.add_paragraph()

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("GOVERNMENT STREET LIGHTING PROJECT")
r.bold = True; r.font.size = Pt(15); r.font.color.rgb = BLUE

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("BUSINESS REQUIREMENTS DOCUMENT")
r.bold = True; r.font.size = Pt(30); r.font.color.rgb = NAVY

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Street Lighting Store, Issuance & Installation\nManagement System  (SLIMS)")
r.bold = True; r.font.size = Pt(18); r.font.color.rgb = BLUE

doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("QR-Based Inventory Control  •  Store & Site Location Management  •  "
              "Digital Requisition & Issuance  •  Geo-Tagged Installation Evidence  •  "
              "Returns of Out-of-Order Items  •  Automated MIS")
r.font.size = Pt(10.5); r.font.color.rgb = GREY

for _ in range(6):
    doc.add_paragraph()

TBL(["Attribute", "Detail"],
    [["Document Title", "Business Requirements Document — SLIMS"],
     ["Project", "Government Street Lighting Project"],
     ["Document Reference", "SLP/BRD/SLIMS/2026/01"],
     ["Version", "1.0 — Baseline for Review"],
     ["Date of Issue", "17 September 2026"],
     ["Classification", "Official — For Internal Circulation"],
     ["Prepared By", "Business Analysis & Systems Team"],
     ["Document Owner", "Project Director, Street Lighting Project"],
     ["Status", "Draft for Stakeholder Sign-Off"]],
    widths=[5.0, 11.0], font=10, hdr_font=10)

doc.add_page_break()

# =========================================================== DOC CONTROL
H1("Document Control")

H2("Version History")
TBL(["Version", "Date", "Author", "Summary of Change", "Status"],
    [["0.1", "02-Sep-2026", "Business Analysis Team", "Initial skeleton and process discovery notes captured from store and field walkthroughs.", "Draft"],
     ["0.5", "09-Sep-2026", "Business Analysis Team", "Functional requirements, business rules and master data model added after supervisor workshops.", "Draft"],
     ["0.9", "15-Sep-2026", "Business Analysis Team", "Geo-tagging rules, returns handling and reporting requirements incorporated.", "Internal Review"],
     ["1.0", "17-Sep-2026", "Business Analysis Team", "Baseline issued for stakeholder review and sign-off. Process flow annexure attached.", "Issued"]],
    widths=[1.6, 2.2, 3.2, 7.0, 2.0])

H2("Reviewers and Approvers")
P("This document is not effective until signed by every approver listed below. Sign-off confirms that the "
  "requirements stated here correctly and completely express the business need, and that they form the "
  "contractual basis for design, build, testing and acceptance.", align="j")
TBL(["Role", "Responsibility", "Name", "Signature", "Date"],
    [["Project Director, Street Lighting Project", "Overall approval and budget authority", "", "", ""],
     ["Deputy Director (Operations)", "Approves store, issuance and field operating model", "", "", ""],
     ["Store In-Charge / Controller of Stores", "Approves inventory, GRN and stock control requirements", "", "", ""],
     ["Lighting Supervisor (Zonal Representative)", "Approves requisition and site-level requirements", "", "", ""],
     ["Manager (Procurement & Contracts)", "Approves PO linkage, vendor and warranty requirements", "", "", ""],
     ["Head of IT / Systems", "Approves technical, integration and security requirements", "", "", ""],
     ["Internal Audit", "Confirms adequacy of controls and audit trail", "", "", ""]],
    widths=[5.2, 5.6, 2.6, 2.4, 2.2])

H2("Distribution List")
TBL(["Recipient", "Purpose"],
    [["Project Director's Office", "Approval and governance"],
     ["Store & Inventory Function (all locations)", "Operational input and implementation readiness"],
     ["Lighting Supervisors and Lighting Officers", "Field process validation and training"],
     ["Procurement & Contracts", "PO, vendor and warranty process alignment"],
     ["IT / Implementation Partner", "Solution design, build and delivery"],
     ["Internal Audit and Finance", "Control, valuation and reconciliation assurance"]],
    widths=[6.5, 9.5])

H2("Related Documents")
TBL(["Ref.", "Document", "Relevance"],
    [["RD-01", "Procurement Rules and Delegation of Financial Powers", "Governs PO issue, approval limits and vendor penalties"],
     ["RD-02", "Store Keeping and Stock Accounting Manual", "Governs GRN, ledgers, physical verification and condemnation"],
     ["RD-03", "Street Lighting Technical Specification Schedule", "Defines article specifications, wattage classes and warranty terms"],
     ["RD-04", "Asset Register and Condemnation Policy", "Governs scrap, disposal and write-off approvals"],
     ["RD-05", "Data Protection and Records Retention Policy", "Governs retention of photographic and location evidence"]],
    widths=[1.4, 7.1, 7.5])

H2("How to Read This Document")
BUL([("FR-xx-nn", "identifies a functional requirement. The two-letter block indicates the module, for example FR-GR-03 is the third requirement of the Goods Receipt module."),
     ("BR-nnn", "identifies a business rule that the system must enforce without exception."),
     ("NFR-nn", "identifies a non-functional requirement covering performance, security, availability and usability."),
     ("Priority", "follows MoSCoW: M (Must have, mandatory for go-live), S (Should have), C (Could have), W (Won't have in this release).")])

# =========================================================== TOC
doc.add_page_break()
H1("Table of Contents")
toc_field()

# =========================================================== 1 EXEC SUMMARY
H1("1.  Executive Summary", page_break=True)

P("The Street Lighting Project procures, stores, issues and installs street lighting material across a large "
  "and geographically dispersed network of sites. Today the chain that connects a purchase order to a "
  "physically installed and working street light is recorded on paper registers and spreadsheets held at "
  "separate locations. Because each link in that chain is recorded independently, the project cannot state "
  "with confidence how much material it holds, where that material physically is, which site consumed it, "
  "whether the item was actually installed at the location it was issued for, or how much of the installed "
  "population has since failed and come back to store.", align="j")

P("This document sets out the business requirements for the Street Lighting Store, Issuance and Installation "
  "Management System (SLIMS): a single system that follows each item of material from vendor delivery to "
  "site installation and, where it fails, back to store. The control mechanism is a unique QR code carried by "
  "every item, created at the point the article is set up against a purchase order and scanned at every "
  "subsequent handover. The verification mechanism is geo-tagged photographic evidence captured before and "
  "after installation, with coordinates stamped onto the image itself and validated by the system against the "
  "geo-fence of the site the material was issued for.", align="j")

P("Five business outcomes justify the investment:", align="j")
BUL([("Physical accountability. ", "Every receipt, issue, installation and return is created by scanning a QR code, so the stock ledger reflects a physical act rather than a later data-entry step."),
     ("Location truth. ", "Material is held against a defined store location down to bin level and issued against a defined site location, so the question 'where is it?' has one answer at every point in time."),
     ("Proof of installation. ", "A street light is only recorded as installed when before-and-after photographs, stamped with coordinates and time, fall inside the geo-fence of the site it was issued for. Installation credit cannot be claimed without it."),
     ("Procurement traceability. ", "Because item codes carry their purchase order reference, consumption, failure rates and warranty claims can be analysed by PO and by vendor."),
     ("Automatic reporting. ", "Activity reports are produced by the system from the transactions themselves, removing the reconciliation effort and the disputes that manual compilation creates.")])

P("The end-to-end process is set out in five phases in Section 8 and illustrated in the cross-functional "
  "process flow at Annexure A. Requirements are specified in Section 10, the rules the system must enforce "
  "in Section 11, and the basis on which the solution will be accepted in Section 21.", align="j")

H2("Summary of Expected Benefits")
TBL(["Benefit", "Measure", "Current Position", "Target Position"],
    [["Stock accuracy", "Variance between system stock and physical count", "Not measurable; reconciliation performed annually at best", "≤ 1% value variance, verified quarterly"],
     ["Issuance traceability", "Percentage of issues traceable to an approved requisition and a named site", "Partial; challan books not consistently linked to sites", "100%"],
     ["Installation verification", "Percentage of installations with geo-validated before/after evidence", "Nil", "≥ 98% (balance handled as documented exceptions)"],
     ["Requisition turnaround", "Requisition raised to material issued", "3–7 working days, paper-driven", "≤ 1 working day where stock is available"],
     ["Reporting effort", "Person-days per month spent compiling activity reports", "Significant manual effort across store and field", "Near zero; reports generated on demand"],
     ["Warranty recovery", "Failed items claimed against vendor warranty within entitlement", "Largely unclaimed; PO linkage not retained", "≥ 90% of eligible failures claimed"]],
    widths=[3.6, 4.2, 4.6, 3.6])

# =========================================================== 2 BACKGROUND
H1("2.  Project Background and Business Context", page_break=True)

H2("2.1  Operating Context")
P("The project maintains street lighting infrastructure across defined administrative zones. Material is "
  "purchased centrally against purchase orders, received into one or more stores, and drawn down by field "
  "teams as installation and maintenance work is scheduled. The work is performed at a very large number of "
  "small, dispersed locations — individual poles, junctions, road stretches and public spaces — which "
  "makes physical supervision of every installation impossible and makes documentary evidence the only "
  "practical control.", align="j")

H2("2.2  Organisational Roles in the Material Chain")
TBL(["Role", "Position in the Chain", "What They Are Accountable For"],
    [["Procurement & Contracts", "Upstream of store", "Raising purchase orders, agreeing specifications, rates and warranty terms, and pursuing vendors on short supply, defect and warranty claims."],
     ["Vendor / Supplier", "Upstream of store", "Delivering material against the purchase order with a delivery challan and invoice, and honouring warranty replacement."],
     ["Store Keeper", "Store", "Physically receiving material, verifying it against the purchase order, labelling and binning it, and issuing it only against an approved requisition."],
     ["Store In-Charge", "Store", "Approving requisitions within delegated limits, controlling stock levels, and certifying physical verification and condemnation."],
     ["Lighting Supervisor", "Demand side", "Planning site work, raising requisitions for the material each site needs, and verifying field exceptions."],
     ["Lighting Officer", "Field", "Taking delivery at site, installing material, capturing before and after evidence, and returning removed or failed items to store."],
     ["PD Office / MIS", "Governance", "Monitoring progress, consumption and exceptions, and reporting to higher authority."],
     ["Internal Audit", "Assurance", "Testing that material purchased was received, issued, installed and accounted for."]],
    widths=[3.4, 3.0, 9.6])

H2("2.3  Why Change Is Needed Now")
P("Three pressures make the current arrangement untenable. First, the volume of installation activity has "
  "grown beyond what paper registers can track, and the reconciliation burden now consumes staff time that "
  "should be spent on field supervision. Second, the project cannot currently demonstrate to audit that "
  "material issued for a site was installed at that site, which exposes it to observation and to loss. "
  "Third, warranty entitlements on failed material are being lost because the link between a failed item and "
  "the purchase order it arrived on is not retained anywhere.", align="j")

# =========================================================== 3 OBJECTIVES
H1("3.  Business Objectives and Success Criteria", page_break=True)

H2("3.1  Business Objectives")
TBL(["Ref.", "Objective", "Rationale"],
    [["OBJ-01", "Establish a single, authoritative record of street lighting material from purchase order to installation, failure and disposal.", "Removes the reconciliation gap between procurement, store and field records."],
     ["OBJ-02", "Make every physical movement of material a scanned event rather than a written entry.", "A scan is contemporaneous with the physical act and cannot be back-dated or estimated."],
     ["OBJ-03", "Hold stock against defined store locations and issue it against defined site locations.", "Gives a definite answer to where material is, at every point in the chain."],
     ["OBJ-04", "Require geo-validated photographic evidence before an installation is accepted as complete.", "Converts an assertion of work done into verifiable evidence."],
     ["OBJ-05", "Retain the purchase order reference on every item code so consumption, failure and warranty can be analysed by PO and vendor.", "Protects warranty entitlement and informs future procurement."],
     ["OBJ-06", "Bring out-of-order material back into the system through a controlled return, repair and condemnation route.", "Closes the loop and prevents failed material from disappearing from the record."],
     ["OBJ-07", "Generate activity, consumption and exception reporting automatically from transactions.", "Removes manual compilation and the disputes it causes."]],
    widths=[1.5, 7.3, 7.2])

H2("3.2  Success Criteria")
P("The project will be judged successful when, three months after go-live, all of the following hold true:", align="j")
NUM(["Every receipt into store, issue out of store, installation at site and return from site in the period is "
     "represented in the system by a scanned transaction with a named user, a timestamp and a location.",
     "A physical count of any two store locations reconciles to the system within a one percent value variance "
     "without adjustment entries being passed to force agreement.",
     "At least ninety-eight percent of installations recorded in the period carry before and after images with "
     "embedded coordinates that fall inside the geo-fence of the site the material was issued for, and every "
     "remaining case carries a documented and approved exception.",
     "Any item code can be traced backwards from its installed site to the requisition, the issue, the GRN, "
     "the purchase order and the vendor, in a single query.",
     "The monthly activity report is produced from the system with no manual compilation, and is accepted by "
     "the PD Office without adjustment.",
     "Warranty claims are raised from the system for failed items still within their warranty period, and the "
     "value of claims raised is reported."])

H2("3.3  Key Performance Indicators")
TBL(["KPI", "Definition", "Target", "Frequency", "Owner"],
    [["Scan compliance", "Transactions created by scan as a share of all stock transactions", "≥ 99%", "Weekly", "Store In-Charge"],
     ["Geo-validation pass rate", "Installations passing geo-fence validation on first submission", "≥ 95%", "Weekly", "Lighting Supervisor"],
     ["Requisition turnaround", "Median hours from requisition submission to issue, where stock exists", "≤ 24 hrs", "Weekly", "Store In-Charge"],
     ["Stock accuracy", "Absolute value variance at physical verification", "≤ 1%", "Quarterly", "Controller of Stores"],
     ["Return compliance", "Removed items returned to store within the prescribed period", "≥ 95%", "Monthly", "Lighting Supervisor"],
     ["Warranty capture", "Eligible failures raised as warranty claims", "≥ 90%", "Monthly", "Procurement"],
     ["Ageing stock", "Value of stock held beyond 180 days without issue", "Declining trend", "Monthly", "Store In-Charge"]],
    widths=[3.0, 6.0, 1.9, 2.1, 3.0])

# =========================================================== 4 SCOPE
H1("4.  Scope", page_break=True)

H2("4.1  In Scope")
TBL(["Area", "Included in This Release"],
    [["Master data", "Store location hierarchy, site location master with geo-fence, article and item code master carrying the PO reference, vendor master, unit of measure, user and role master."],
     ["QR / barcode", "Generation of unique codes at article creation, label printing, and scanning at receipt, issue, installation and return."],
     ["Inward", "Goods receipt against purchase order by scan, discrepancy and short-receipt handling, binning to store location, stock ledger posting."],
     ["Requisition", "Digital requisition raised by the Lighting Supervisor against a site location, availability check, approval workflow, rejection and revision."],
     ["Issuance", "Scan-based issue against an approved requisition, stock reservation, gate pass and issue note, transfer of custody to the site location."],
     ["Installation", "Mobile capture of before and after photographs with coordinates and time stamped onto the image, scan of the item at site, geo-fence validation, installation confirmation and asset tagging."],
     ["Returns", "Return of out-of-order and dismantled items to store by scan, quarantine, repair and warranty routing, condemnation and scrap register."],
     ["Inventory control", "Stock ledger, reservations, ageing, re-order level alerts, inter-store transfer, physical verification and adjustment."],
     ["Reporting", "Activity, consumption, installation progress, geo-map, exception, ageing, out-of-order and warranty reporting with export."],
     ["Access control", "Role-based permissions, delegated approval limits and a complete audit trail."]],
    widths=[3.2, 12.8])

H2("4.2  Out of Scope")
TBL(["Excluded Area", "Reason / Treatment"],
    [["Tendering, bid evaluation and purchase order creation", "Performed in the existing procurement process. SLIMS consumes the approved PO as an input."],
     ["Vendor payment, invoice processing and financial accounting entries", "Remain with the finance function. SLIMS supplies GRN and consumption data through an interface."],
     ["Payroll, attendance and human resource management", "No overlap with the material chain."],
     ["Energy billing, metering and electricity consumption analysis", "A separate operational domain; may be considered in a later phase."],
     ["Smart or IoT-enabled lighting telemetry and remote fault detection", "Deferred. The data model will not prevent a later integration."],
     ["Civil, electrical and contractor work billing and measurement books", "Handled under contract management outside this system."],
     ["Mobile network provision and field hardware procurement", "An organisational dependency, not a system requirement."]],
    widths=[7.0, 9.0])

H2("4.3  Assumptions Governing Scope")
BUL(["Purchase orders are approved outside SLIMS and are made available to it, either by interface or by "
     "controlled entry, before material is received against them.",
     "Every site at which material may be installed can be defined as a site location with coordinates before "
     "material is issued to it.",
     "Field staff are equipped with smartphones capable of camera, GPS and QR scanning, and are enrolled as "
     "named users of the system.",
     "Physical QR labels are durable enough to survive outdoor installation and remain scannable for the "
     "expected service life of the item."])

# =========================================================== 5 STAKEHOLDERS
H1("5.  Stakeholders and Responsibility Matrix", page_break=True)

H2("5.1  Stakeholder Register")
TBL(["Stakeholder", "Interest in the System", "Influence", "Engagement Approach"],
    [["Project Director", "Assurance that material purchased translates into installed, working lights", "High", "Approves the BRD; receives monthly MIS; chairs steering review"],
     ["Deputy Director (Operations)", "Field productivity and installation progress against plan", "High", "Weekly progress review; validates the field operating model"],
     ["Controller of Stores / Store In-Charge", "Stock accuracy, accountability and clean physical verification", "High", "Process design workshops; owns store UAT"],
     ["Store Keeper", "A workable receipt and issue routine that is faster than the register", "Medium", "Usability testing; training; day-one floor support"],
     ["Lighting Supervisor", "Getting the right material to the right site on time", "High", "Requisition design; owns site master data; field UAT"],
     ["Lighting Officer", "A mobile app that works at site, including where the network is weak", "Medium", "Field pilot; offline testing; training"],
     ["Procurement & Contracts", "PO linkage, vendor performance and warranty recovery", "Medium", "Interface design; warranty workflow validation"],
     ["Finance", "Valuation, consumption posting and asset capitalisation", "Medium", "Interface and valuation rule sign-off"],
     ["Internal Audit", "Sufficiency of controls and the completeness of the audit trail", "Medium", "Control walkthrough before go-live"],
     ["IT / Implementation Partner", "Buildable, supportable requirements", "High", "Owns design, build, deployment and support"]],
    widths=[3.4, 6.0, 1.8, 4.8])

H2("5.2  RACI for the Core Process")
P("R = Responsible, A = Accountable, C = Consulted, I = Informed.", italic=True, size=9)
TBL(["Process Step", "Store Keeper", "Store In-Charge", "Lighting Supervisor", "Lighting Officer", "Procurement", "PD Office"],
    [["Article and item code creation with PO reference", "C", "A", "I", "I", "R", "I"],
     ["QR label generation and printing", "R", "A", "I", "I", "C", "I"],
     ["Goods receipt and GRN posting", "R", "A", "I", "I", "C", "I"],
     ["Discrepancy and short-receipt handling", "R", "C", "I", "I", "A", "I"],
     ["Binning to store location", "R", "A", "I", "I", "I", "I"],
     ["Raising the digital requisition", "I", "C", "R/A", "C", "I", "I"],
     ["Requisition approval", "I", "R/A", "C", "I", "I", "I"],
     ["Issue against requisition by scan", "R", "A", "I", "C", "I", "I"],
     ["Receipt of material at site", "I", "I", "C", "R/A", "I", "I"],
     ["Before and after evidence capture", "I", "I", "C", "R/A", "I", "I"],
     ["Geo-fence exception verification", "I", "I", "R/A", "C", "I", "I"],
     ["Return of out-of-order items", "R", "C", "C", "R", "I", "I"],
     ["Repair, warranty and condemnation routing", "R", "A", "I", "I", "R", "I"],
     ["Physical verification and adjustment", "R", "A", "I", "I", "I", "C"],
     ["MIS review and escalation", "I", "C", "C", "I", "C", "R/A"]],
    widths=[5.0, 1.9, 2.0, 2.2, 2.0, 1.6, 1.5], font=8.5, hdr_font=8.5,
    align_center=[1, 2, 3, 4, 5, 6])

# =========================================================== 6 AS-IS
H1("6.  Current State and Pain Points", page_break=True)

H2("6.1  How the Process Works Today")
NUM(["Procurement issues a purchase order. The store is informed informally or by copy of the order.",
     "The vendor delivers material with a delivery challan. The Store Keeper counts it and writes it into the "
     "bin card and the stock register. Any shortage is noted on the challan and pursued by telephone.",
     "Material is stacked wherever space allows. Location is retained in the Store Keeper's memory rather than "
     "in any record.",
     "A Lighting Supervisor who needs material sends a written requisition, or telephones the store. Approval "
     "is a signature on the slip.",
     "The Store Keeper issues against the slip, writes a gate pass, and records the issue in the register. The "
     "site the material is destined for may or may not be recorded.",
     "The Lighting Officer installs the material. Completion is reported verbally or in a periodic written "
     "return. No independent evidence of where the work was done is captured.",
     "Failed material removed from site is sometimes brought back to store and stacked in a corner; often it "
     "is left at site or discarded. Warranty is rarely claimed.",
     "At month end, activity figures are compiled by hand from several registers, and rarely agree with each "
     "other."])

H2("6.2  Pain Points and Consequences")
TBL(["#", "Pain Point", "Business Consequence", "How SLIMS Addresses It"],
    [["P1", "No physical location recorded against stock", "Material cannot be found; issues are delayed; duplicate purchases are made against stock already held", "Store location hierarchy to bin level, mandatory at GRN (FR-GR-07)"],
     ["P2", "Receipt quantities recorded by manual count without PO validation", "Short and excess receipts are absorbed silently; vendor penalties are not applied", "Scan-based receipt validated against PO line quantity (FR-GR-03, BR-014)"],
     ["P3", "Requisitions on paper, with no site reference", "Consumption cannot be attributed to a site; demand cannot be planned", "Digital requisition mandatorily tied to a site location (FR-RQ-02)"],
     ["P4", "Issues recorded without item-level identity", "The individual item cannot be traced after it leaves the store", "Every item carries a unique QR identity from creation (FR-MD-06)"],
     ["P5", "No evidence that installation happened at the intended site", "Installation cannot be verified; audit observations and potential leakage", "Geo-stamped before/after images validated against the site geo-fence (FR-IN-04 to FR-IN-08)"],
     ["P6", "Failed items not returned or recorded", "Warranty entitlement is lost; failure patterns are invisible", "Controlled return, quarantine and warranty routing (FR-RT-01 to FR-RT-07)"],
     ["P7", "PO reference not retained after receipt", "Vendor and batch performance cannot be analysed", "PO reference carried on the item code permanently (BR-004)"],
     ["P8", "Reports compiled manually from disagreeing registers", "Management decisions taken on unreliable figures", "Reports generated from the transaction record itself (FR-RP-01)"],
     ["P9", "No user attribution on stock movements", "Accountability cannot be established after the event", "Every transaction stamped with the named user, device and time (NFR-13)"]],
    widths=[0.9, 4.2, 5.2, 5.7])

# =========================================================== 7 TO-BE
H1("7.  Proposed Solution Overview", page_break=True)

H2("7.1  Solution Concept")
P("SLIMS treats the QR code as the identity of a physical item and the geo-tagged photograph as the proof of "
  "a physical act. Every transaction in the system is the record of something that happened in the physical "
  "world, captured at the moment it happened, by a named user, at a known location. The system does not "
  "accept a stock movement that was not scanned, and it does not accept an installation that was not "
  "photographed inside the geo-fence of the site the material was issued for.", align="j")

H2("7.2  Solution Components")
TBL(["Component", "Users", "Purpose"],
    [["Master data and configuration module", "Store In-Charge, IT administrator", "Defines store locations, site locations and geo-fences, articles and item codes with PO reference, vendors, users, roles and approval limits."],
     ["QR generation and label service", "Store In-Charge, Store Keeper", "Creates a unique, non-reusable code for each item or batch at article creation and produces printable, weather-resistant labels."],
     ["Store desktop application", "Store Keeper, Store In-Charge", "Goods receipt, binning, requisition approval, issue, returns, quarantine, physical verification and store reporting."],
     ["Supervisor portal", "Lighting Supervisor", "Site master maintenance, digital requisition, exception verification, site-wise progress and consumption view."],
     ["Field mobile application", "Lighting Officer", "Acknowledge receipt at site, capture geo-stamped before and after images, scan the item at the point of installation, record returns. Works offline and synchronises when connectivity returns."],
     ["Rules and validation engine", "System", "Enforces the business rules in Section 11, including PO matching, stock reservation, geo-fence tolerance and status transitions."],
     ["Reporting and MIS", "PD Office, Audit, all functions", "Pre-built activity, consumption, progress, geo-map, exception, ageing and warranty reports with drill-down and export."],
     ["Audit log", "Internal Audit, IT", "Immutable record of every transaction, override and master data change."]],
    widths=[4.0, 3.6, 8.4])

H2("7.3  Guiding Design Principles")
BUL([("Scan or it did not happen. ", "Stock quantity changes only through a scanned transaction. Manual quantity entry is an exception requiring a reason code and a higher approval."),
     ("Location is mandatory. ", "Stock always sits at a store location; issued material always belongs to a site location. There is no unassigned state."),
     ("Evidence travels with the claim. ", "An installation claim is inseparable from the images and coordinates that support it; neither can be edited after submission."),
     ("The purchase order reference never falls off. ", "It is carried on the item code from creation through installation, failure and disposal."),
     ("Exceptions are recorded, not prevented into oblivion. ", "Where the field cannot comply, the system records a reasoned, approved exception rather than allowing a silent workaround."),
     ("The field works offline. ", "No field action may depend on live connectivity; the application queues and synchronises.")])

# =========================================================== 8 PHASES
H1("8.  The To-Be Process in Five Phases", page_break=True)
P("The process illustrated at Annexure A is described below. Each phase names the trigger, the actor, the "
  "steps and the control that makes the phase reliable.", align="j")

H2("8.1  Phase 1 — Procurement Inward and Goods Receipt")
P("Trigger: a purchase order is approved and material is despatched by the vendor.", italic=True, size=9.5)
NUM(["The article is created in the item master, article-wise, with its item code, specification, unit of "
     "measure and rate, and is mapped to the purchase order number and line it was bought against.",
     "The system generates a unique QR code for each item or batch to be received, embedding the item code, "
     "the PO reference and a serial number, and produces printable labels.",
     "The vendor delivers the consignment to the nominated store with a delivery challan and invoice.",
     "The Store Keeper affixes the label where the vendor has not pre-labelled, and scans each item or batch "
     "at the inward gate. The system matches every scan against the open PO line.",
     "If quantity, specification or PO reference do not match, the system blocks the receipt of the affected "
     "line and generates a discrepancy or short-receipt note for Procurement to pursue with the vendor.",
     "On a clean match, the GRN is posted and the Store Keeper bins the material to a specific store, rack "
     "and bin location.",
     "The stock ledger is updated store-location-wise, PO-wise and item-code-wise, and the GRN register "
     "becomes visible to the PD Office."])
P("Control: nothing enters stock without a scan that matches an open purchase order line, and nothing is "
  "binned without a location.", size=9.5, italic=True)

H2("8.2  Phase 2 — Digital Requisition and Approval")
P("Trigger: a Lighting Supervisor plans work at one or more sites.", italic=True, size=9.5)
NUM(["The Supervisor raises a digital requisition in the system, selecting the site location, the item codes "
     "required, the quantity of each and the date the material is needed. Free-text site descriptions are not "
     "accepted; the site must already exist in the site master with coordinates.",
     "The system checks availability at the nominated store. Where stock is short, it raises an automatic "
     "purchase indent to Procurement for the shortfall and tells the Supervisor what can and cannot be met.",
     "The requisition is routed for approval to the Store In-Charge, or to a higher authority where the value "
     "exceeds the delegated limit.",
     "A rejected requisition returns to the Supervisor with a reason and may be revised and resubmitted.",
     "On approval, the system reserves the stock against the requisition so that it cannot be issued "
     "elsewhere, and generates an Issue Slip."])
P("Control: demand is captured against a real site before material moves, and approved stock is ring-fenced.",
  size=9.5, italic=True)

H2("8.3  Phase 3 — Issuance to the Site Location")
P("Trigger: an approved requisition with stock reserved.", italic=True, size=9.5)
NUM(["The Store Keeper calls up the approved requisition and scans the QR code of each item being issued.",
     "The system validates each scan against three things at once: that the item code matches a requisition "
     "line, that the quantity does not exceed what was approved, and that the item is not already issued, "
     "installed or quarantined. An invalid scan is rejected on the spot with the reason.",
     "On a valid scan set, the issue is posted. Custody of each item moves from the store location to the "
     "site location named on the requisition, and the system generates a gate pass and issue note.",
     "The Lighting Officer acknowledges receipt of the material at site on the mobile application."])
P("Control: material can only leave the store against an approved requisition, and only to the site that "
  "requisition names.", size=9.5, italic=True)

H2("8.4  Phase 4 — Installation with Geo-Tagged Evidence")
P("Trigger: material acknowledged at site.", italic=True, size=9.5)
NUM(["Before beginning work, the Lighting Officer captures a BEFORE photograph of the installation point. "
     "The application stamps the latitude, longitude, date and time onto the image itself at the moment of "
     "capture, and records the same values as data alongside it.",
     "The Officer installs the fixture and scans the item's QR code at the pole or site point, confirming that "
     "this specific item was installed here.",
     "The Officer captures an AFTER photograph, stamped in the same way.",
     "The system compares the coordinates embedded in the evidence against the geo-fence of the site location "
     "the item was issued for, applying the configured tolerance.",
     "Where the coordinates fall outside the geo-fence, the system does not accept the installation. It raises "
     "an exception to the Lighting Supervisor for field verification, and the Officer may be required to "
     "recapture the evidence.",
     "Where the coordinates fall inside the geo-fence, the installation is confirmed. The item is tagged to "
     "the site as an asset, its status becomes INSTALLED, and the consumption is posted.",
     "The activity report for the site, the officer and the period is updated automatically."])
P("Control: installation credit is evidence-based. The photograph carries the coordinates, and the system, "
  "not the officer, decides whether those coordinates are acceptable.", size=9.5, italic=True)

H2("8.5  Phase 5 — Return of Out-of-Order Items, Repair and Condemnation")
P("Trigger: an item fails in service, or is dismantled during replacement work.", italic=True, size=9.5)
NUM(["The Lighting Officer records the removal against the site, states the reason for failure, and brings "
     "the item back to the store.",
     "The Store Keeper scans the returned item. The system recognises it, recovers its full history including "
     "its purchase order and installation date, and posts a Return Note into a quarantine location.",
     "The item is classified as repairable, warranty-eligible or fit only for scrap.",
     "A warranty-eligible item is routed to Procurement for a claim against the vendor. A repairable item goes "
     "for repair. In either case the restored or replaced item re-enters stock through the GRN cycle and is "
     "available for issue again.",
     "An item fit only for scrap is moved to the condemnation bin, put forward for disposal approval, and on "
     "approval written into the scrap register and out of stock.",
     "Stock, asset and scrap ledgers and the site history are updated, and failure analysis by article, vendor "
     "and purchase order becomes available."])
P("Control: failed material cannot disappear. It is scanned back in, its warranty entitlement is tested, and "
  "its disposal is approved and recorded.", size=9.5, italic=True)
